import io
from datetime import date, datetime
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from .domain import HISTORY_END, HISTORY_START, ROUTES
from .storage import SnapshotStore, atomic_write, digest, file_hash


def parquet_write(path: Path, frame: pd.DataFrame):
    buf = io.BytesIO()
    frame.to_parquet(buf, index=False)
    atomic_write(path, buf.getvalue())


def build_history(labels: pd.DataFrame) -> pd.DataFrame:
    required = {"route", "date", "hour", "boardings"}
    if set(labels.columns) != required:
        raise ValueError(f"Expected columns {sorted(required)}")
    labels = labels.copy()
    for column in ("route", "hour", "boardings"):
        if not labels[column].astype(str).str.fullmatch(r"\d+").all():
            raise ValueError(f"Invalid nonnegative integer in {column}")
        labels[column] = labels[column].astype("int64")
    if not labels["date"].astype(str).str.fullmatch(r"\d{4}-\d{2}-\d{2}").all():
        raise ValueError("Expected ISO local dates")
    if not labels.route.isin(ROUTES).all() or not labels.hour.between(0, 23).all():
        raise ValueError("Unknown route or invalid hour")
    labels["timestamp"] = pd.to_datetime(labels["date"], format="%Y-%m-%d").dt.tz_localize(
        "Europe/Moscow"
    ) + pd.to_timedelta(labels.hour, unit="h")
    if not ((labels.timestamp >= HISTORY_START) & (labels.timestamp < HISTORY_END)).all():
        raise ValueError("Labels outside the verified history window")
    if labels.duplicated(["route", "timestamp"]).any():
        raise ValueError("Duplicate route-hour label key")
    hours = pd.date_range(HISTORY_START, HISTORY_END, freq="h", inclusive="left")
    grid = pd.MultiIndex.from_product([ROUTES, hours], names=["route", "timestamp"])
    values = labels.set_index(["route", "timestamp"]).boardings.reindex(grid)
    frame = values.rename("value").reset_index()
    frame["valueOrigin"] = np.where(frame.value.isna(), "filled_zero", "label")
    frame["value"] = frame.value.fillna(0).astype("int64")
    frame["coverageStatus"] = "provided_extract"
    frame["qualityFlag"] = np.where(
        (frame.route == 50) & frame.timestamp.dt.strftime("%Y-%m-%d").isin(["2025-09-20", "2025-09-21"]),
        "requires_review",
        "none",
    )
    frame["allOperations"] = None
    frame["failedOperations"] = None
    groups = frame.groupby([frame.route, frame.timestamp.dt.dayofweek, frame.timestamp.dt.hour]).value
    frame["baseline"] = groups.transform(lambda s: s.shift().rolling(8, min_periods=1).mean())
    frame["baselineCount"] = groups.transform(lambda s: s.shift().rolling(8, min_periods=1).count()).fillna(0)
    return frame


def prepare_labels(store: SnapshotStore, dataset: Path):
    names = ["labels/labels_day_train.csv", "labels/labels_day_test.csv"]
    inputs = {name: file_hash(dataset / name) for name in names}
    ident = "history-" + digest({"inputs": inputs, "schema": 2, "timezone": "Europe/Moscow"})
    if store.path("histories", ident).exists():
        current = store.current()
        store.publish(
            historyId=ident,
            forecastId=current.get("forecastId") if current.get("historyId") == ident else None,
            fleetId=current.get("fleetId") if current.get("historyId") == ident else None,
        )
        return store.read("histories", ident)
    parts = [pd.read_csv(dataset / name, sep=";", dtype=str, keep_default_na=False) for name in names]
    frame = build_history(pd.concat(parts, ignore_index=True))
    # These expectations identify this competition dataset, not arbitrary future feeds.
    expected_hashes = [
        "20a61009b4ee00327784c6fa888109609fd3a1905ea3a8509b36964dd9baced7",
        "37479acea3e319ac613d72df7217bbe57b66af49644e538139174fd534ce2347",
    ]
    known = list(inputs.values()) == expected_hashes
    checks = {
        "keys": len(frame),
        "positiveKeys": int((frame.value > 0).sum()),
        "filledZeros": int((frame.valueOrigin == "filled_zero").sum()),
        "sum": int(frame.value.sum()),
        "knownDataset": known,
        "rawAuditExecuted": False,
    }
    if known and (checks["keys"], checks["positiveKeys"], checks["filledZeros"], checks["sum"]) != (
        72960,
        57551,
        15409,
        59667191,
    ):
        raise ValueError("Competition dataset invariant failed")
    if known and [int(pd.to_numeric(p.boardings).sum()) for p in parts] != [46912710, 12754481]:
        raise ValueError("Train/test label totals differ")
    manifest = {
        "id": ident,
        "schemaVersion": 2,
        "timezone": "Europe/Moscow",
        "inputs": inputs,
        "start": HISTORY_START.isoformat(),
        "end": HISTORY_END.isoformat(),
        "checks": checks,
        "metric": "successful_validations",
        "coverage": "Completeness of the extract is not proven",
        "rawAudit": {
            "executed": False,
            "reference": "analysis/dataset/reconciliation.json",
            "reportedHistoricalDifferences": 0 if known else None,
        },
    }
    parquet_write(store.path("histories", ident, "parquet"), frame)
    store.put("histories", ident, manifest)
    current = store.current()
    store.publish(
        historyId=ident,
        forecastId=current.get("forecastId") if current.get("historyId") == ident else None,
        fleetId=current.get("fleetId") if current.get("historyId") == ident else None,
    )
    return manifest


def _iso(value):
    if value is None or value == "":
        return None
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m-%d")
    return pd.to_datetime(value, dayfirst=False).strftime("%Y-%m-%d")


def _id(value):
    return str(int(value)) if isinstance(value, (int, float)) else str(value).strip()


def prepare_network(store: SnapshotStore, dataset: Path):
    from .geography import GEOGRAPHY_FILE, prepare_osm_network
    from .network_history import HISTORY_INDEX, prepare_history

    if (dataset / "geography" / HISTORY_INDEX).is_file():
        return prepare_history(store, dataset / "geography")
    if (dataset / "geography/history-2025").exists():
        raise ValueError("Geography archive is incomplete: run scripts/fetch_geography_history.py")

    mapped_source = dataset / "geography" / GEOGRAPHY_FILE
    if mapped_source.is_file():
        return prepare_osm_network(store, mapped_source, historical=True)
    source = dataset / "spravochniki/Хакатон_справочники_трамвай_10_маршрутов.xlsx"
    sha = file_hash(source)
    ident = "network-" + digest({"sha256": sha, "schema": 2})
    if store.path("networks", ident).exists():
        store.publish(networkId=ident)
        return store.read("networks", ident)
    wb = load_workbook(source, read_only=True, data_only=True)

    def table(name):
        rows = iter(wb[name].iter_rows(values_only=True))
        next(rows)  # Human-readable explanation row.
        header = next(rows)
        return [dict(zip(header, row)) for row in rows if any(v is not None for v in row)]

    route_rows = table("Маршруты GTFS_ROUTES")
    stop_rows = table("Остановки GTFS_STOPS")
    visits = table("Порядок_остановок GTFS_TRIPS_ST")
    wb.close()
    routes_by_id = {_id(r["route_id"]): r for r in route_rows}
    stations = {_id(s["stop_id"]): s for s in stop_rows}
    patterns, stops, segments = [], [], []
    for trip in sorted({_id(v["trip_id"]) for v in visits}):
        rows = sorted([r for r in visits if _id(r["trip_id"]) == trip], key=lambda r: int(r["stop_sequence"]))
        route = routes_by_id[_id(rows[0]["route_id"])]
        route_id = _id(route["route_short_name"])
        pid = f"pattern-{trip}"
        seq = [int(r["stop_sequence"]) for r in rows]
        if seq != list(range(1, len(rows) + 1)):
            raise ValueError("Non-contiguous or duplicate stop occurrences")
        actual = [_iso(route["actual_date"])]
        points = []
        for row in rows:
            station = stations[_id(row["stop_id"])]
            coordinate = [float(station["stop_lon"]), float(station["stop_lat"])]
            if not (-180 <= coordinate[0] <= 180 and -90 <= coordinate[1] <= 90):
                raise ValueError("Invalid network coordinates")
            actual.extend([_iso(row["actual_date"]), _iso(station["actual_date"])])
            visit = {
                "id": f"occ-{trip}-{row['stop_sequence']}",
                "stationId": _id(row["stop_id"]),
                "name": str(station["stop_name"]),
                "routeId": route_id,
                "patternId": pid,
                "direction": int(row["direction_id"]),
                "order": int(row["stop_sequence"]),
                "coordinates": coordinate,
            }
            points.append(visit)
            stops.append(visit)
        for first, second in zip(points, points[1:]):
            segments.append(
                {
                    "id": f"segment-{trip}-{first['order']}",
                    "routeId": route_id,
                    "patternId": pid,
                    "fromId": first["id"],
                    "toId": second["id"],
                    "order": first["order"],
                    "coordinates": [first["coordinates"], second["coordinates"]],
                }
            )
        valid_from = max(filter(None, [_iso(rows[0]["start_date"]), _iso(route["route_date_start"])]))
        valid_ends = list(filter(None, [_iso(rows[0]["end_date"]), _iso(route["route_date_end"])]))
        patterns.append(
            {
                "id": pid,
                "routeId": route_id,
                "direction": int(rows[0]["direction_id"]),
                "name": f"{points[0]['name']} → {points[-1]['name']}",
                "validFrom": valid_from,
                "validTo": min(valid_ends) if valid_ends else None,
                "observedAt": max(filter(None, actual)),
            }
        )
    geo_routes = {_id(r["route_short_name"]): r for r in route_rows}
    routes = []
    for route_id in sorted(set(geo_routes) | {str(r) for r in ROUTES}, key=int):
        source_route = geo_routes.get(route_id)
        routes.append(
            {
                "id": route_id,
                "number": route_id,
                "name": source_route["route_long_name"] if source_route else f"Маршрут № {route_id}",
                "isTarget": int(route_id) in ROUTES,
                "hasData": int(route_id) in ROUTES,
                "hasGeometry": source_route is not None,
                "historySupport": "no_positive_history"
                if route_id == "5"
                else ("available" if int(route_id) in ROUTES else "unavailable"),
            }
        )
    network = {
        "networkSnapshotId": ident,
        "source": source.name,
        "sourceHash": sha,
        "geometryQuality": "schematic",
        "asOf": max(p["observedAt"] for p in patterns),
        "routes": routes,
        "patterns": patterns,
        "stops": stops,
        "segments": segments,
        "missingGeometryRouteIds": [str(r) for r in ROUTES if str(r) not in geo_routes],
        "warning": "Справочная схема по позднему срезу. Историческая применимость ограничена датами источника.",
    }
    store.put("networks", ident, network)
    store.publish(networkId=ident)
    return network
